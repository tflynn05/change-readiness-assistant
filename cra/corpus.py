"""Load documents from the public and private knowledge base folders.

Public documents are markdown files with provenance front matter written by
scripts/build_corpus.py. Private documents are whatever you put in
corpus/private/ or in a folder named by CRA_PRIVATE_DIR: markdown, text, CSV,
PDF, Word, PowerPoint, or Excel (the last four need `pip install -e ".[files]"`).
Private files are read in place and never copied into the project.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path

import yaml

FRONT_MATTER = re.compile(r"\A---\s*\n(.*?)\n---\s*\n", re.S)
PRIVATE_LICENSE = "Private: not for redistribution"
TEXT_TYPES = {".md", ".markdown", ".txt"}
SUPPORTED = TEXT_TYPES | {".csv", ".pdf", ".docx", ".pptx", ".xlsx"}


@dataclass(frozen=True)
class Document:
    doc_id: str          # "<source_id>/<file stem>", stable across rebuilds
    title: str
    text: str
    source_id: str
    source_title: str
    source_url: str
    license: str
    private: bool = False


def _front_matter(raw: str) -> tuple[dict, str]:
    m = FRONT_MATTER.match(raw)
    if not m:
        return {}, raw
    meta = yaml.safe_load(m.group(1)) or {}
    return (meta if isinstance(meta, dict) else {}), raw[m.end():]


def _first_heading(text: str) -> str | None:
    m = re.search(r"^#{1,2}\s+(.+)$", text, flags=re.M)
    return m.group(1).strip() if m else None


def _pretty_stem(path: Path) -> str:
    return re.sub(r"[_\s]+", " ", path.stem).strip()


def _need(module: str, package: str, path: Path):
    try:
        return __import__(module)
    except ImportError as exc:
        raise RuntimeError(f"{path.name}: install {package} to index this file "
                           f"(pip install -e \".[files]\")") from exc


def _cell(value) -> str:
    return " ".join(str(value).split()).replace("|", "/") if value is not None else ""


def _table(rows: list[list[str]]) -> str:
    """Rows to a markdown table, dropping empty rows and empty trailing columns."""
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return ""
    width = max(max((i + 1 for i, c in enumerate(r) if c), default=0) for r in rows)
    rows = [(r + [""] * width)[:width] for r in rows]
    header, body = rows[0], rows[1:]
    lines = ["| " + " | ".join(header) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


def read_pdf(path: Path) -> str:
    import logging
    pypdf = _need("pypdf", "pypdf", path)
    logging.getLogger("pypdf").setLevel(logging.ERROR)   # font-encoding notices are noise here
    reader = pypdf.PdfReader(str(path))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    # Each page becomes a section, so a citation points to a page.
    return "\n\n".join(f"## Page {i}\n\n{text}" for i, text in enumerate(pages, 1) if text)


def read_docx(path: Path) -> str:
    docx = _need("docx", "python-docx", path)
    document = docx.Document(str(path))
    out = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "2"
            out.append("#" * min(int(level) + 1, 4) + " " + text)
        else:
            out.append(text)
    for table in document.tables:
        out.append(_table([[_cell(c.text) for c in row.cells] for row in table.rows]))
    return "\n\n".join(o for o in out if o)


def read_pptx(path: Path) -> str:
    _need("pptx", "python-pptx", path)
    from pptx import Presentation
    out = []
    for n, slide in enumerate(Presentation(str(path)).slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        parts = []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(r.text for r in para.runs).strip()
                    if text:
                        parts.append(("- " if para.level else "") + text)
            if getattr(shape, "has_table", False) and shape.has_table:
                parts.append(_table([[_cell(c.text) for c in row.cells] for row in shape.table.rows]))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes: {notes}")
        if title or parts:
            out.append(f"## Slide {n}" + (f": {title}" if title else "") + "\n\n" + "\n\n".join(parts))
    return "\n\n".join(out)


def read_xlsx(path: Path) -> str:
    openpyxl = _need("openpyxl", "openpyxl", path)
    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out = []
    for sheet in workbook.worksheets:
        rows = [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        table = _table(rows)
        if table:
            out.append(f"## Sheet: {sheet.title}\n\n{table}")
    workbook.close()
    return "\n\n".join(out)


def read_csv(path: Path) -> str:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    return _table([[_cell(c) for c in row] for row in csv.reader(io.StringIO(text))])


READERS = {".pdf": read_pdf, ".docx": read_docx, ".pptx": read_pptx, ".xlsx": read_xlsx, ".csv": read_csv}


def load_public(directory: Path) -> list[Document]:
    docs = []
    for path in sorted(directory.rglob("*.md")):
        meta, body = _front_matter(path.read_text(encoding="utf-8"))
        source_id = meta.get("source_id") or path.parent.name
        docs.append(Document(
            doc_id=f"{source_id}/{path.stem}",
            title=str(meta.get("title") or _first_heading(body) or path.stem),
            text=body,
            source_id=source_id,
            source_title=str(meta.get("source_title") or source_id),
            source_url=str(meta.get("source_url") or path.as_posix()),
            license=str(meta.get("license") or "unknown"),
        ))
    return docs


def private_files(directory: Path) -> list[Path]:
    if not directory.exists():
        return []
    return sorted(p for p in directory.rglob("*") if p.is_file()
                  and p.suffix.lower() in SUPPORTED
                  and not p.name.startswith((".", "~$"))
                  and p.name.lower() != "readme.md")


def load_private(directory: Path, prefix: str = "", errors: list | None = None) -> list[Document]:
    """Read every supported file under `directory`. Files that fail to parse are skipped
    and reported in `errors` rather than stopping the index build."""
    docs = []
    for path in private_files(directory):
        rel = (Path(prefix) / path.relative_to(directory)).as_posix()
        suffix = path.suffix.lower()
        try:
            if suffix in TEXT_TYPES:
                meta, body = _front_matter(path.read_text(encoding="utf-8", errors="replace"))
            else:
                meta, body = {}, READERS[suffix](path)
        except Exception as exc:  # one bad file should not block the rest
            if errors is not None:
                errors.append(f"{rel}: {exc}")
            continue
        if not body.strip():
            if errors is not None:
                errors.append(f"{rel}: no extractable text")
            continue
        heading = _first_heading(body) if suffix in TEXT_TYPES | {".docx"} else None
        docs.append(Document(
            doc_id=f"private/{rel.rsplit('.', 1)[0]}",
            title=str(meta.get("title") or heading or _pretty_stem(path)),
            text=body,
            source_id="private",
            source_title=str(meta.get("source_title") or "Private knowledge base"),
            source_url=f"private:{rel}",
            license=PRIVATE_LICENSE,
            private=True,
        ))
    return docs


def load_documents(public_dir: Path, private_dirs: list[Path] | None = None,
                   errors: list | None = None) -> list[Document]:
    docs = load_public(public_dir)
    for i, directory in enumerate(private_dirs or []):
        prefix = "" if i == 0 else directory.name     # first is corpus/private; others keep their folder name
        docs += load_private(directory, prefix, errors)
    return docs
